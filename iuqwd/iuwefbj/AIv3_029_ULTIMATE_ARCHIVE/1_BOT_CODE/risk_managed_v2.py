"""
End-to-end test: 100 bots → top 10 → risk-managed top 10 → ensemble.

1. Generate 100 bots deterministically (seed=42)
2. Run all 100 on last 30 days WITHOUT risk management → baseline
3. Pick actual top 10 by net profit
4. Run those exact 10 WITH 7-layer risk management + proper R:R + trailing stops
5. Build ensemble of top 5 risk-managed bots
6. Write Excel with everything
"""
from __future__ import annotations

import json
import time
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from backtester import Costs, ema, rsi, atr, adx, bollinger, vwap_rolling, donchian
from smart_strategies import zscore, resample_htf, taker_buy_ratio, reindex_htf_to_ltf
from smart_strategies_round2 import htf_strong_uptrend, htf_strong_downtrend
from bots_100 import (
    trend_filter, entry_trigger, confirmation, stop_price,
    generate_100_strategies, rolling_linreg_slope,
    INITIAL_WALLET, RISK_PER_TRADE, COSTS, DATA_DIR, DOWNLOAD,
)

END_DATE = pd.Timestamp("2026-10-14 18:15:00", tz="UTC")
START_DATE = END_DATE - pd.Timedelta(days=30)


def load_window(symbol: str, timeframe: str):
    if timeframe == "4h":
        df_1h = pd.read_parquet(DATA_DIR / f"{symbol}_1h.parquet").set_index("open_time").sort_index()
        df = df_1h.resample("4h").agg({
            "open": "first", "high": "max", "low": "min",
            "close": "last", "volume": "sum", "quote_volume": "sum",
            "taker_buy_base": "sum",
        }).dropna()
    else:
        df = pd.read_parquet(DATA_DIR / f"{symbol}_{timeframe}.parquet").set_index("open_time").sort_index()
    warmup = START_DATE - pd.Timedelta(hours=400)
    df = df.loc[warmup:END_DATE].copy()
    df_1h_raw = pd.read_parquet(DATA_DIR / f"{symbol}_1h.parquet").set_index("open_time").sort_index()
    htf = df_1h_raw.loc[warmup:END_DATE].copy()
    return df, htf


# ============================================================================
# BASELINE backtester (no risk mgmt) — same as bots_100.backtest_strategy
# ============================================================================

def backtest_baseline(df, htf, cfg):
    trend_long, trend_short = trend_filter(df, htf, cfg["trend"])
    confirm_mask = confirmation(df, cfg["confirm"])
    if cfg["side"] == "long":
        sig_long = entry_trigger(df, cfg["entry"], "long") & trend_long & confirm_mask
        sig_short = pd.Series(False, index=df.index)
        stop_long = stop_price(df, cfg["stop"], "long")
        stop_short = pd.Series(np.nan, index=df.index)
    elif cfg["side"] == "short":
        sig_short = entry_trigger(df, cfg["entry"], "short") & trend_short & confirm_mask
        sig_long = pd.Series(False, index=df.index)
        stop_short = stop_price(df, cfg["stop"], "short")
        stop_long = pd.Series(np.nan, index=df.index)
    else:
        sig_long = entry_trigger(df, cfg["entry"], "long") & trend_long & confirm_mask
        sig_short = entry_trigger(df, cfg["entry"], "short") & trend_short & confirm_mask
        stop_long = stop_price(df, cfg["stop"], "long")
        stop_short = stop_price(df, cfg["stop"], "short")
    r_mult = cfg["target"]; max_bars = cfg["max_bars"]

    df = df.sort_index(); n = len(df)
    if n < 50: return [], INITIAL_WALLET

    sig_long_arr = sig_long.fillna(False).values
    sig_short_arr = sig_short.fillna(False).values
    stop_long_arr = stop_long.values
    stop_short_arr = stop_short.values
    opens = df["open"].values; highs = df["high"].values
    lows = df["low"].values; closes = df["close"].values
    times = df.index
    slip_frac = COSTS.slippage_bps / 10000.0

    trades = []; wallet = INITIAL_WALLET; i = 0
    while i < n:
        signal_long = i+1 < n and sig_long_arr[i] and np.isfinite(opens[i+1]) and np.isfinite(stop_long_arr[i])
        signal_short = i+1 < n and sig_short_arr[i] and np.isfinite(opens[i+1]) and np.isfinite(stop_short_arr[i])
        chosen = None
        if signal_long and signal_short: chosen = "long"
        elif signal_long: chosen = "long"
        elif signal_short: chosen = "short"
        if chosen is None: i += 1; continue
        entry_raw = float(opens[i+1])
        if entry_raw <= 0: i += 1; continue
        if chosen == "long":
            sp = float(stop_long_arr[i])
            if not np.isfinite(sp) or sp >= entry_raw: i += 1; continue
            entry_price = entry_raw * (1 + slip_frac)
            target_price = entry_price + r_mult * (entry_price - sp)
        else:
            sp = float(stop_short_arr[i])
            if not np.isfinite(sp) or sp <= entry_raw: i += 1; continue
            entry_price = entry_raw * (1 - slip_frac)
            target_price = entry_price - r_mult * (sp - entry_price)
        risk_per_unit = abs(entry_price - sp)
        if risk_per_unit <= 0: i += 1; continue
        risk_amount = RISK_PER_TRADE * wallet
        size_units = risk_amount / risk_per_unit
        if size_units <= 0: i += 1; continue
        entry_notional = size_units * entry_price
        entry_commission = entry_notional * COSTS.commission_pct
        wallet_after_entry = wallet - entry_commission
        exit_idx = -1; exit_price = float("nan"); exit_reason = ""; bars_held = 0
        j = i + 1
        while j < n:
            bars_held = j - i
            h = highs[j]; l = lows[j]
            if chosen == "long":
                hit_stop = l <= sp; hit_tgt = h >= target_price
            else:
                hit_stop = h >= sp; hit_tgt = l <= target_price
            if hit_stop and hit_tgt:
                exit_price = sp; exit_reason = "stop"; exit_idx = j; break
            elif hit_stop:
                exit_price = sp; exit_reason = "stop"; exit_idx = j; break
            elif hit_tgt:
                exit_price = target_price; exit_reason = "target"; exit_idx = j; break
            if bars_held >= max_bars:
                exit_price = closes[j]; exit_reason = "time"; exit_idx = j; break
            j += 1
        if exit_idx < 0:
            exit_price = closes[-1]; exit_reason = "end_of_data"; exit_idx = n - 1
        if chosen == "long":
            exit_fill = exit_price * (1 - slip_frac)
            gross_pnl = size_units * (exit_fill - entry_price)
        else:
            exit_fill = exit_price * (1 + slip_frac)
            gross_pnl = size_units * (entry_price - exit_fill)
        exit_notional = size_units * exit_fill
        exit_commission = exit_notional * COSTS.commission_pct
        total_cost = entry_commission + exit_commission
        net_pnl = gross_pnl - total_cost
        r_mult_actual = net_pnl / risk_amount if risk_amount > 0 else 0.0
        trades.append({
            "entry_time": times[i+1], "exit_time": times[exit_idx],
            "side": chosen, "entry_price": entry_price, "exit_price": exit_fill,
            "stop": sp, "target": target_price, "size": size_units,
            "risk_pct": RISK_PER_TRADE, "risk_usd": risk_amount,
            "gross_pnl": gross_pnl, "cost": total_cost, "net_pnl": net_pnl,
            "wallet_before": wallet, "wallet_after": wallet_after_entry + net_pnl,
            "R_multiple": r_mult_actual, "exit_reason": exit_reason, "bars_held": bars_held,
        })
        wallet = wallet_after_entry + net_pnl
        i = exit_idx + 1
    return trades, wallet


# ============================================================================
# RISK-MANAGED backtester — 7 layers + proper R:R + trailing stops
# ============================================================================

@dataclass
class RiskConfig:
    base_risk_pct: float = 0.01
    monthly_loss_cap_pct: float = 0.08
    weekly_loss_cap_pct: float = 0.05
    daily_loss_cap_pct: float = 0.03
    dd_throttle_1: float = 0.10
    dd_throttle_2: float = 0.15
    dd_throttle_3: float = 0.20
    consec_loss_throttle_1: int = 3
    consec_loss_throttle_2: int = 5
    consec_loss_throttle_1_trades: int = 2
    consec_loss_throttle_2_hours: int = 24
    anti_mart_win_mult: float = 1.20
    anti_mart_loss_mult: float = 0.70
    anti_mart_max_risk: float = 0.015
    anti_mart_min_risk: float = 0.004
    end_of_month_days: int = 3
    end_of_month_risk: float = 0.004
    target_r_mult: float = 2.0       # FIXED at 1:2 R:R
    enable_trailing: bool = True
    trail_trigger_1: float = 1.0
    trail_trigger_2: float = 1.5
    trail_distance_r: float = 0.5
    max_bars: int = 24


def backtest_risk_managed(df, htf, cfg, risk_cfg=RiskConfig()):
    """Run strategy with 7-layer risk management + proper R:R + trailing stops.
    NOTE: target_r_mult is OVERRIDDEN to 2.0 (proper R:R fix per spec)."""
    trend_long, trend_short = trend_filter(df, htf, cfg["trend"])
    confirm_mask = confirmation(df, cfg["confirm"])
    if cfg["side"] == "long":
        sig_long = entry_trigger(df, cfg["entry"], "long") & trend_long & confirm_mask
        sig_short = pd.Series(False, index=df.index)
        stop_long = stop_price(df, cfg["stop"], "long")
        stop_short = pd.Series(np.nan, index=df.index)
    elif cfg["side"] == "short":
        sig_short = entry_trigger(df, cfg["entry"], "short") & trend_short & confirm_mask
        sig_long = pd.Series(False, index=df.index)
        stop_short = stop_price(df, cfg["stop"], "short")
        stop_long = pd.Series(np.nan, index=df.index)
    else:
        sig_long = entry_trigger(df, cfg["entry"], "long") & trend_long & confirm_mask
        sig_short = entry_trigger(df, cfg["entry"], "short") & trend_short & confirm_mask
        stop_long = stop_price(df, cfg["stop"], "long")
        stop_short = stop_price(df, cfg["stop"], "short")

    r_mult = risk_cfg.target_r_mult  # FIXED at 2.0
    max_bars = risk_cfg.max_bars

    df = df.sort_index(); n = len(df)
    if n < 50: return [], {}

    sig_long_arr = sig_long.fillna(False).values
    sig_short_arr = sig_short.fillna(False).values
    stop_long_arr = stop_long.values
    stop_short_arr = stop_short.values
    opens = df["open"].values; highs = df["high"].values
    lows = df["low"].values; closes = df["close"].values
    times = df.index
    slip_frac = COSTS.slippage_bps / 10000.0

    # State
    wallet = INITIAL_WALLET
    peak_wallet = INITIAL_WALLET
    base_risk = risk_cfg.base_risk_pct
    cur_risk = base_risk
    consec_losses = 0; consec_wins = 0
    throttle_remaining = 0
    pause_until = None
    monthly_start = INITIAL_WALLET; weekly_start = INITIAL_WALLET; daily_start = INITIAL_WALLET
    cur_month = times[0].month; cur_week = times[0].isocalendar()[1]; cur_day = times[0].day
    monthly_stopped = False; weekly_stopped = False; daily_stopped = False; dd_stopped = False

    trades = []; skip_reasons = {}
    i = 0
    while i < n:
        t = times[i]
        # Reset counters on date rollovers
        if t.day != cur_day:
            daily_stopped = False; daily_start = wallet; cur_day = t.day
        if t.isocalendar()[1] != cur_week:
            weekly_stopped = False; weekly_start = wallet; cur_week = t.isocalendar()[1]
        if t.month != cur_month:
            monthly_stopped = False; monthly_start = wallet; cur_month = t.month

        if wallet > peak_wallet: peak_wallet = wallet

        # Check all 7 risk layers
        allowed = True; reason = "ok"; eff_risk = cur_risk
        if monthly_stopped: allowed = False; reason = "monthly_cap"
        elif weekly_stopped: allowed = False; reason = "weekly_cap"
        elif daily_stopped: allowed = False; reason = "daily_cap"
        elif dd_stopped: allowed = False; reason = "dd_cap"
        elif pause_until is not None and t < pause_until: allowed = False; reason = "consec_pause"

        if allowed:
            # Check P&L caps
            if monthly_start > 0:
                mp = (wallet - monthly_start) / monthly_start
                if mp <= -risk_cfg.monthly_loss_cap_pct:
                    monthly_stopped = True; allowed = False; reason = "monthly_cap_hit"
            if allowed and weekly_start > 0:
                wp = (wallet - weekly_start) / weekly_start
                if wp <= -risk_cfg.weekly_loss_cap_pct:
                    weekly_stopped = True; allowed = False; reason = "weekly_cap_hit"
            if allowed and daily_start > 0:
                dp = (wallet - daily_start) / daily_start
                if dp <= -risk_cfg.daily_loss_cap_pct:
                    daily_stopped = True; allowed = False; reason = "daily_cap_hit"

        if allowed:
            # DD throttle
            if peak_wallet > 0:
                dd = (peak_wallet - wallet) / peak_wallet
                if dd >= risk_cfg.dd_throttle_3:
                    dd_stopped = True; allowed = False; reason = "dd_cap_hit"
                else:
                    dd_stopped = False
                    if dd >= risk_cfg.dd_throttle_2:
                        eff_risk = min(eff_risk, base_risk * 0.25)
                    elif dd >= risk_cfg.dd_throttle_1:
                        eff_risk = min(eff_risk, base_risk * 0.50)

        if allowed:
            # End-of-month protection
            if t.month == 12:
                next_month = pd.Timestamp(year=t.year+1, month=1, day=1, tz="UTC")
            else:
                next_month = pd.Timestamp(year=t.year, month=t.month+1, day=1, tz="UTC")
            days_left = (next_month - t).days
            if days_left <= risk_cfg.end_of_month_days and monthly_start > 0:
                mp = (wallet - monthly_start) / monthly_start
                if mp > 0:
                    eff_risk = risk_cfg.end_of_month_risk
            # Throttle remaining
            if throttle_remaining > 0:
                eff_risk = max(eff_risk * 0.5, risk_cfg.anti_mart_min_risk)
                throttle_remaining -= 1

        if not allowed:
            skip_reasons[reason] = skip_reasons.get(reason, 0) + 1
            i += 1; continue

        # Check for signal
        signal_long = i+1 < n and sig_long_arr[i] and np.isfinite(opens[i+1]) and np.isfinite(stop_long_arr[i])
        signal_short = i+1 < n and sig_short_arr[i] and np.isfinite(opens[i+1]) and np.isfinite(stop_short_arr[i])
        chosen = None
        if signal_long and signal_short: chosen = "long"
        elif signal_long: chosen = "long"
        elif signal_short: chosen = "short"
        if chosen is None: i += 1; continue

        # === PROPER R:R FIX: stop from fill price ===
        entry_raw = float(opens[i+1])
        if entry_raw <= 0: i += 1; continue
        if chosen == "long":
            stop_ref = float(stop_long_arr[i])
            if not np.isfinite(stop_ref) or stop_ref >= entry_raw: i += 1; continue
            entry_price = entry_raw * (1 + slip_frac)
            stop_distance = abs(entry_raw - stop_ref)
            stop_price_val = entry_price - stop_distance
            target_price = entry_price + r_mult * (entry_price - stop_price_val)
        else:
            stop_ref = float(stop_short_arr[i])
            if not np.isfinite(stop_ref) or stop_ref <= entry_raw: i += 1; continue
            entry_price = entry_raw * (1 - slip_frac)
            stop_distance = abs(stop_ref - entry_raw)
            stop_price_val = entry_price + stop_distance
            target_price = entry_price - r_mult * (stop_price_val - entry_price)
        if stop_distance <= 0: i += 1; continue

        risk_amount = eff_risk * wallet
        size_units = risk_amount / stop_distance
        if size_units <= 0: i += 1; continue

        entry_notional = size_units * entry_price
        entry_commission = entry_notional * COSTS.commission_pct
        wallet_after_entry = wallet - entry_commission

        # === Walk forward with trailing stop ===
        exit_idx = -1; exit_price = float("nan"); exit_reason = ""; bars_held = 0
        cur_trail_stop = stop_price_val
        best_price = entry_price
        j = i + 1
        while j < n:
            bars_held = j - i
            h = highs[j]; l = lows[j]
            if chosen == "long":
                if h > best_price: best_price = h
            else:
                if l < best_price: best_price = l
            if risk_cfg.enable_trailing:
                r_best = (best_price - entry_price) / stop_distance if chosen == "long" \
                         else (entry_price - best_price) / stop_distance
                if r_best >= risk_cfg.trail_trigger_1:
                    if chosen == "long":
                        if entry_price > cur_trail_stop: cur_trail_stop = entry_price
                    else:
                        if entry_price < cur_trail_stop: cur_trail_stop = entry_price
                if r_best >= risk_cfg.trail_trigger_2:
                    if chosen == "long":
                        new_stop = best_price - risk_cfg.trail_distance_r * stop_distance
                        if new_stop > cur_trail_stop: cur_trail_stop = new_stop
                    else:
                        new_stop = best_price + risk_cfg.trail_distance_r * stop_distance
                        if new_stop < cur_trail_stop: cur_trail_stop = new_stop
            if chosen == "long":
                hit_stop = l <= cur_trail_stop; hit_tgt = h >= target_price
            else:
                hit_stop = h >= cur_trail_stop; hit_tgt = l <= target_price
            if hit_stop and hit_tgt:
                exit_price = cur_trail_stop; exit_reason = "stop"; exit_idx = j; break
            elif hit_stop:
                if cur_trail_stop >= entry_price: exit_reason = "trail_win"
                else: exit_reason = "stop"
                exit_price = cur_trail_stop; exit_idx = j; break
            elif hit_tgt:
                exit_price = target_price; exit_reason = "target"; exit_idx = j; break
            if bars_held >= max_bars:
                exit_price = closes[j]; exit_reason = "time"; exit_idx = j; break
            j += 1
        if exit_idx < 0:
            exit_price = closes[-1]; exit_reason = "end_of_data"; exit_idx = n - 1
        if chosen == "long":
            exit_fill = exit_price * (1 - slip_frac)
            gross_pnl = size_units * (exit_fill - entry_price)
        else:
            exit_fill = exit_price * (1 + slip_frac)
            gross_pnl = size_units * (entry_price - exit_fill)
        exit_notional = size_units * exit_fill
        exit_commission = exit_notional * COSTS.commission_pct
        total_cost = entry_commission + exit_commission
        net_pnl = gross_pnl - total_cost
        r_mult_actual = net_pnl / risk_amount if risk_amount > 0 else 0.0

        # Update state
        new_wallet = wallet_after_entry + net_pnl
        won = net_pnl > 0
        if won:
            consec_wins += 1; consec_losses = 0
            cur_risk = min(cur_risk * risk_cfg.anti_mart_win_mult, risk_cfg.anti_mart_max_risk)
        else:
            consec_losses += 1; consec_wins = 0
            cur_risk = max(cur_risk * risk_cfg.anti_mart_loss_mult, risk_cfg.anti_mart_min_risk)
        # Consecutive loss throttle
        if consec_losses >= risk_cfg.consec_loss_throttle_2:
            pause_until = times[exit_idx] + pd.Timedelta(hours=risk_cfg.consec_loss_throttle_2_hours)
            consec_losses = 0
        elif consec_losses >= risk_cfg.consec_loss_throttle_1:
            throttle_remaining = risk_cfg.consec_loss_throttle_1_trades
            cur_risk = max(cur_risk * 0.5, risk_cfg.anti_mart_min_risk)

        trades.append({
            "entry_time": times[i+1], "exit_time": times[exit_idx],
            "side": chosen, "entry_price": entry_price, "exit_price": exit_fill,
            "stop": stop_price_val, "target": target_price, "size": size_units,
            "risk_pct": eff_risk, "risk_usd": risk_amount,
            "gross_pnl": gross_pnl, "cost": total_cost, "net_pnl": net_pnl,
            "wallet_before": wallet, "wallet_after": new_wallet,
            "R_multiple": r_mult_actual, "exit_reason": exit_reason, "bars_held": bars_held,
        })
        wallet = new_wallet
        if wallet > peak_wallet: peak_wallet = wallet
        i = exit_idx + 1

    # Stats
    if trades:
        wallets = [INITIAL_WALLET] + [t["wallet_after"] for t in trades]
        peak_arr = np.maximum.accumulate(wallets)
        dd_arr = (peak_arr - wallets) / np.where(peak_arr > 0, peak_arr, 1)
        max_dd = float(np.max(dd_arr))
    else:
        max_dd = 0
    stats = {
        "n_trades": len(trades), "final_wallet": wallet,
        "net": wallet - INITIAL_WALLET, "net_pct": (wallet - INITIAL_WALLET) / INITIAL_WALLET,
        "peak_wallet": peak_wallet, "max_dd_pct": max_dd, "skip_reasons": skip_reasons,
    }
    return trades, stats


def main():
    print("="*120)
    print("END-TO-END TEST: 100 bots → top 10 → risk-managed top 10 → ensemble")
    print(f"Window: {START_DATE} → {END_DATE}  |  Initial: ${INITIAL_WALLET:,.2f}")
    print("="*120)

    # Step 1: Generate 100 bots
    configs = generate_100_strategies(seed=42)
    print(f"\nStep 1: Generated {len(configs)} bots")

    # Step 2: Load data once
    print("\nStep 2: Loading data...")
    data_cache = {}
    for sym in ["BTCUSDT", "ETHUSDT"]:
        for tf in ["1h", "4h"]:
            data_cache[(sym, tf)] = load_window(sym, tf)
            print(f"  {sym} {tf}: {len(data_cache[(sym,tf)][0])} bars")

    # Step 3: Run all 100 baseline
    print("\nStep 3: Running all 100 bots BASELINE (no risk mgmt)...")
    t0 = time.time()
    all_baseline = []
    for cfg in configs:
        df, htf = data_cache[(cfg["symbol"], cfg["timeframe"])]
        try:
            trades, final = backtest_baseline(df, htf, cfg)
            net = final - INITIAL_WALLET
            n_wins = sum(1 for t in trades if t["exit_reason"] == "target" and t["net_pnl"] > 0)
            n_losses = sum(1 for t in trades if t["exit_reason"] == "stop" and t["net_pnl"] < 0)
            wr = n_wins/(n_wins+n_losses) if (n_wins+n_losses) > 0 else 0
            all_baseline.append({"cfg": cfg, "trades": trades, "final": final, "net": net, "wr": wr, "n_trades": len(trades)})
        except Exception as e:
            all_baseline.append({"cfg": cfg, "trades": [], "final": INITIAL_WALLET, "net": 0, "wr": 0, "n_trades": 0})
    print(f"  Done in {time.time()-t0:.1f}s")

    # Step 4: Pick top 10 by net profit
    all_baseline.sort(key=lambda x: x["net"], reverse=True)
    top_10 = all_baseline[:10]
    print(f"\nStep 4: Top 10 by baseline net profit:")
    for rank, r in enumerate(top_10, 1):
        cfg = r["cfg"]
        print(f"  #{rank} {cfg['id'][:55]:<55} {cfg['symbol']:<8} {cfg['timeframe']} {cfg['side']:<5} {r['n_trades']:>3}t WR={r['wr']*100:>5.1f}% net=${r['net']:>+9,.2f} ({r['net']/INITIAL_WALLET*100:>+6.2f}%)")

    # Step 5: Run top 10 with risk management
    print(f"\nStep 5: Running top 10 WITH risk management + proper R:R (fixed 2.0) + trailing stops...")
    risk_cfg = RiskConfig()
    rm_results = []
    for r in top_10:
        cfg = r["cfg"]
        df, htf = data_cache[(cfg["symbol"], cfg["timeframe"])]
        trades_rm, stats_rm = backtest_risk_managed(df, htf, cfg, risk_cfg)
        n_wins_rm = sum(1 for t in trades_rm if t["net_pnl"] > 0)
        n_losses_rm = sum(1 for t in trades_rm if t["net_pnl"] < 0)
        wr_rm = n_wins_rm/(n_wins_rm+n_losses_rm) if (n_wins_rm+n_losses_rm) > 0 else 0
        rm_results.append({
            "cfg": cfg, "trades_base": r["trades"], "net_base": r["net"], "wr_base": r["wr"], "n_trades_base": r["n_trades"],
            "trades_rm": trades_rm, "net_rm": stats_rm["net"], "net_pct_rm": stats_rm["net_pct"],
            "final_rm": stats_rm["final_wallet"], "wr_rm": wr_rm, "n_trades_rm": stats_rm["n_trades"],
            "max_dd_rm": stats_rm["max_dd_pct"], "skip_reasons": stats_rm["skip_reasons"],
            "peak_wallet": stats_rm["peak_wallet"],
        })

    # Print comparison
    print("\n" + "="*130)
    print("COMPARISON: BASELINE vs RISK-MANAGED (top 10)")
    print("="*130)
    print(f"{'#':<3} {'Bot':<10} {'Sym':<8} {'TF':<4} {'Side':<5} | {'Base trades':>11} {'Base WR':>7} {'Base net':>10} | {'RM trades':>9} {'RM WR':>6} {'RM net':>10} {'RM DD':>7} | {'Δ Net':>9}")
    print("-"*130)
    for rank, r in enumerate(rm_results, 1):
        cfg = r["cfg"]
        delta = r["net_rm"] - r["net_base"]
        print(f"{rank:<3} {cfg['id'][:10]:<10} {cfg['symbol']:<8} {cfg['timeframe']:<4} {cfg['side']:<5} | "
              f"{r['n_trades_base']:>11} {r['wr_base']*100:>6.1f}% {r['net_base']:>+10,.2f} | "
              f"{r['n_trades_rm']:>9} {r['wr_rm']*100:>5.1f}% {r['net_rm']:>+10,.2f} {r['max_dd_rm']*100:>6.1f}% | "
              f"{delta:>+9,.2f}")

    # Aggregate stats
    base_total = sum(r["net_base"] for r in rm_results)
    rm_total = sum(r["net_rm"] for r in rm_results)
    base_wins = sum(1 for r in rm_results if r["net_base"] > 0)
    rm_wins = sum(1 for r in rm_results if r["net_rm"] > 0)
    base_avg_dd = np.mean([r["max_dd_rm"] for r in rm_results])  # only have RM DD
    print(f"\n  Total net (sum 10 bots):  baseline ${base_total:>+,.2f}  →  RM ${rm_total:>+,.2f}  (Δ ${rm_total-base_total:>+,.2f})")
    print(f"  # profitable:             baseline {base_wins}/10  →  RM {rm_wins}/10")
    print(f"  Avg max DD (RM):          {base_avg_dd*100:.1f}%")

    # Top 5 ensemble
    rm_results.sort(key=lambda x: x["net_rm"], reverse=True)
    print(f"\n--- ENSEMBLE: $2,000 each on top 5 risk-managed bots ---")
    top5 = rm_results[:5]
    ens_total = sum(r["net_rm"] * 0.2 for r in top5)
    print(f"  Initial: ${INITIAL_WALLET:,.2f}  →  Final: ${INITIAL_WALLET + ens_total:,.2f}")
    print(f"  Net: ${ens_total:+,.2f} ({ens_total/INITIAL_WALLET*100:+.2f}%)")
    print(f"  Avg max DD: {np.mean([r['max_dd_rm'] for r in top5])*100:.1f}%")
    for r in top5:
        cfg = r["cfg"]
        print(f"    {cfg['id'][:10]:10s} ({cfg['symbol']} {cfg['timeframe']} {cfg['side']}): "
              f"${r['net_rm']*0.2:+9,.2f} on $2,000  ({r['net_pct_rm']*100:+6.2f}%)  DD={r['max_dd_rm']*100:.1f}%")

    # Write Excel
    out_path = DOWNLOAD / "risk_managed_30day_v2.xlsx"
    write_excel(rm_results, out_path)
    print(f"\n>>> Wrote {out_path}")


def write_excel(rm_results, out_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    pos_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    neg_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    warn_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    border = Border(left=Side(style='thin'), right=Side(style='thin'),
                    top=Side(style='thin'), bottom=Side(style='thin'))

    ws = wb.create_sheet("Summary")
    ws["A1"] = "RISK-MANAGED 30-DAY TEST — Top 10 bots, baseline vs 7-layer risk mgmt + proper R:R + trailing stops"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
    ws["A2"] = f"Window: {START_DATE} → {END_DATE}  |  Initial: ${INITIAL_WALLET:,.2f}  |  Base risk: 1%  |  Fixed R:R = 1:2"
    ws["A3"] = "Risk: monthly -8% / weekly -5% / daily -3% caps, DD throttle 10/15/20%, consec-loss throttle 3/5, anti-martingale ±20/30%, EOM protection 0.4%"

    headers = ["Rank", "Bot ID", "Symbol", "TF", "Side", "Trend", "Entry", "Stop", "Orig Target",
               "Base Trades", "Base WR", "Base Net $", "Base Net %",
               "RM Trades", "RM WR", "RM Net $", "RM Net %", "RM Max DD %",
               "Δ Net $ (RM-Base)", "Skip Reasons"]
    row_idx = 5
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row_idx, column=col, value=h)
        c.font = header_font; c.fill = header_fill; c.alignment = Alignment(horizontal="center", wrap_text=True); c.border = border
    for rank, r in enumerate(rm_results, 1):
        cfg = r["cfg"]
        delta = r["net_rm"] - r["net_base"]
        row = [rank, cfg["id"], cfg["symbol"], cfg["timeframe"], cfg["side"],
               cfg["trend"], cfg["entry"], cfg["stop"], cfg["target"],
               r["n_trades_base"], r["wr_base"], r["net_base"], r["net_base"]/INITIAL_WALLET,
               r["n_trades_rm"], r["wr_rm"], r["net_rm"], r["net_pct_rm"], r["max_dd_rm"],
               delta, json.dumps(r["skip_reasons"], default=str)]
        for col, v in enumerate(row, 1):
            c = ws.cell(row=row_idx+rank, column=col, value=v)
            c.border = border
            if col in (12, 16, 19):
                c.number_format = '"$"#,##0.00'
            elif col in (13, 17, 18):
                c.number_format = '0.00%'
            elif col in (11, 15):
                c.number_format = '0.0%'
            if col == 16 and isinstance(v, (int, float)):
                c.fill = pos_fill if v > 0 else (neg_fill if v < 0 else warn_fill)
            if col == 19 and isinstance(v, (int, float)):
                c.fill = pos_fill if v > 0 else (neg_fill if v < 0 else warn_fill)
    widths = [6, 10, 9, 5, 6, 14, 22, 14, 9, 9, 9, 12, 10, 9, 9, 12, 10, 10, 14, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A6"

    # Per-bot trade logs (top 10)
    for rank, r in enumerate(rm_results, 1):
        cfg = r["cfg"]; trades = r["trades_rm"]
        sheet_name = f"#{rank} {cfg['id']}"[:31]
        ws = wb.create_sheet(sheet_name)
        ws["A1"] = f"#{rank} — {cfg['id']}  |  {cfg['symbol']} {cfg['timeframe']} {cfg['side']}  |  RM Net: ${r['net_rm']:+,.2f} ({r['net_pct_rm']*100:+.2f}%)  |  Trades: {len(trades)}  |  Max DD: {r['max_dd_rm']*100:.1f}%"
        ws["A1"].font = Font(bold=True, size=11); ws.merge_cells("A1:S1")
        ws["A2"] = f"Trend={cfg['trend']}  Entry={cfg['entry']}  Stop={cfg['stop']}  Target=2.0R (FIXED)  |  Skips: {r['skip_reasons']}"
        ws.merge_cells("A2:S2")
        trade_headers = ["#", "Entry Time", "Exit Time", "Side", "Entry $", "Exit $",
                         "Stop $", "Target $", "Size", "Risk %", "Risk $", "Gross P&L $",
                         "Cost $", "Net P&L $", "R Mult", "Exit Reason", "Bars Held",
                         "Wallet Before $", "Wallet After $"]
        for col, h in enumerate(trade_headers, 1):
            c = ws.cell(row=4, column=col, value=h)
            c.font = header_font; c.fill = header_fill; c.alignment = Alignment(horizontal="center"); c.border = border
        for i, t in enumerate(trades, 1):
            entry_t = t["entry_time"].tz_localize(None) if hasattr(t["entry_time"], "tz_localize") else t["entry_time"]
            exit_t = t["exit_time"].tz_localize(None) if hasattr(t["exit_time"], "tz_localize") else t["exit_time"]
            row = [i, entry_t, exit_t, t["side"], t["entry_price"], t["exit_price"],
                   t["stop"], t["target"], t["size"], t["risk_pct"], t["risk_usd"],
                   t["gross_pnl"], t["cost"], t["net_pnl"], t["R_multiple"],
                   t["exit_reason"], t["bars_held"], t["wallet_before"], t["wallet_after"]]
            for col, v in enumerate(row, 1):
                c = ws.cell(row=4+i, column=col, value=v)
                c.border = border
                if col in (5, 6, 7, 8, 11, 12, 13, 14, 18, 19) and isinstance(v, (int, float)):
                    c.number_format = '"$"#,##0.00'
                elif col == 10 and isinstance(v, (int, float)):
                    c.number_format = '0.00%'
                elif col == 15 and isinstance(v, (int, float)):
                    c.number_format = '0.00'
                if col == 14 and isinstance(v, (int, float)):
                    c.fill = pos_fill if v > 0 else (neg_fill if v < 0 else warn_fill)
        twidths = [5, 22, 22, 6, 12, 12, 12, 12, 12, 8, 10, 12, 10, 12, 8, 12, 8, 14, 14]
        for i, w in enumerate(twidths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A5"
    wb.save(out_path)


if __name__ == "__main__":
    main()
